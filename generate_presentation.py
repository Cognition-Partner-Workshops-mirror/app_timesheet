"""
Generate AI FinOps Market Landscape & Whitespace Analysis PowerPoint presentation.
Uses python-pptx to create a professional 15-slide deck with light blue color palette.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE

# ---------------------------------------------------------------------------
# Color palette
# ---------------------------------------------------------------------------
LIGHTEST_BLUE = RGBColor(0xE3, 0xF2, 0xFD)   # #E3F2FD
LIGHT_BLUE    = RGBColor(0xBB, 0xDE, 0xFB)    # #BBDEFB
MID_BLUE      = RGBColor(0x90, 0xCA, 0xF9)    # #90CAF9
BLUE          = RGBColor(0x64, 0xB5, 0xF6)     # #64B5F6
ACCENT_BLUE   = RGBColor(0x42, 0xA5, 0xF5)    # #42A5F5
DARK_BLUE     = RGBColor(0x1E, 0x88, 0xE5)    # #1E88E5
NAVY          = RGBColor(0x1A, 0x23, 0x7E)     # #1A237E
DARK_GRAY     = RGBColor(0x33, 0x33, 0x33)     # #333333
WHITE         = RGBColor(0xFF, 0xFF, 0xFF)
VERY_LIGHT    = RGBColor(0xF5, 0xF5, 0xF5)

FONT_NAME = "Calibri"

SLIDE_WIDTH  = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

prs = Presentation()
prs.slide_width  = SLIDE_WIDTH
prs.slide_height = SLIDE_HEIGHT


# ---------------------------------------------------------------------------
# Helper functions
# ---------------------------------------------------------------------------
def _set_slide_bg(slide, color=WHITE):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def _add_circle(slide, left, top, size, fill_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape


def _add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def _set_text(tf, text, font_size=14, color=DARK_GRAY, bold=False, alignment=PP_ALIGN.LEFT, font_name=FONT_NAME):
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return p


def _add_paragraph(tf, text, font_size=14, color=DARK_GRAY, bold=False, alignment=PP_ALIGN.LEFT,
                   space_before=Pt(4), space_after=Pt(2), bullet=False, level=0, font_name=FONT_NAME):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    p.level = level
    if space_before is not None:
        p.space_before = space_before
    if space_after is not None:
        p.space_after = space_after
    return p


def _add_slide_title(slide, title_text, left=Inches(0.6), top=Inches(0.3),
                     width=Inches(12), height=Inches(0.7)):
    """Add a styled slide title with accent bar."""
    # Accent bar
    _add_rect(slide, Inches(0.6), Inches(0.25), Inches(0.12), Inches(0.55), DARK_BLUE)
    tb = _add_textbox(slide, Inches(0.9), top, width, height)
    _set_text(tb.text_frame, title_text, font_size=28, color=DARK_BLUE, bold=True)
    # Divider line below title
    _add_rect(slide, Inches(0.6), Inches(1.0), Inches(12.1), Pt(2), LIGHT_BLUE)


def _build_table(slide, left, top, width, col_widths, headers, rows, row_height=Inches(0.45)):
    """Create a styled table on the slide."""
    num_rows = len(rows) + 1
    num_cols = len(headers)
    table_height = row_height * num_rows

    tbl_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, table_height)
    table = tbl_shape.table

    # Set column widths
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    # Header row
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT_BLUE
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(12)
            paragraph.font.color.rgb = WHITE
            paragraph.font.bold = True
            paragraph.font.name = FONT_NAME
            paragraph.alignment = PP_ALIGN.LEFT
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Data rows
    for r_idx, row_data in enumerate(rows):
        bg = WHITE if r_idx % 2 == 0 else LIGHTEST_BLUE
        for c_idx, val in enumerate(row_data):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            # Bold the solution name (first column)
            if c_idx == 0:
                p.font.bold = True
                p.font.color.rgb = NAVY
            else:
                p.font.bold = False
                p.font.color.rgb = DARK_GRAY
            p.text = val
            p.font.size = Pt(10)
            p.font.name = FONT_NAME
            cell.vertical_anchor = MSO_ANCHOR.TOP

    return tbl_shape


# ===========================================================================
# SLIDE 1 : Title Slide
# ===========================================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
_set_slide_bg(slide1, WHITE)

# Large blue banner at top
_add_rect(slide1, Inches(0), Inches(0), SLIDE_WIDTH, Inches(3.8), DARK_BLUE)
# Decorative lighter strip
_add_rect(slide1, Inches(0), Inches(3.8), SLIDE_WIDTH, Inches(0.15), ACCENT_BLUE)
# Small circle accents
_add_circle(slide1, Inches(11.5), Inches(0.5), Inches(0.9), MID_BLUE)
_add_circle(slide1, Inches(12.0), Inches(1.5), Inches(0.5), BLUE)

# Title
tb = _add_textbox(slide1, Inches(1), Inches(1.0), Inches(11), Inches(1.2))
_set_text(tb.text_frame, "AI FinOps: Market Landscape & Whitespace Analysis",
          font_size=36, color=WHITE, bold=True, alignment=PP_ALIGN.LEFT)

# Subtitle
tb2 = _add_textbox(slide1, Inches(1), Inches(2.3), Inches(11), Inches(0.6))
_set_text(tb2.text_frame, "Telemetry  \u00b7  Governance  \u00b7  Metrics  \u00b7  Insights",
          font_size=20, color=LIGHTEST_BLUE, bold=False, alignment=PP_ALIGN.LEFT)

# Date
tb3 = _add_textbox(slide1, Inches(1), Inches(3.1), Inches(5), Inches(0.5))
_set_text(tb3.text_frame, "June 2026", font_size=16, color=LIGHT_BLUE, bold=False)

# Decorative bottom accent rectangles
_add_rect(slide1, Inches(1), Inches(5.0), Inches(1.8), Inches(0.08), MID_BLUE)
_add_rect(slide1, Inches(3.0), Inches(5.0), Inches(1.0), Inches(0.08), ACCENT_BLUE)


# ===========================================================================
# SLIDE 2 : Executive Summary
# ===========================================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
_set_slide_bg(slide2, WHITE)
_add_slide_title(slide2, "Executive Summary")

bullets = [
    "AI spending is growing exponentially \u2014 LLM API costs, GPU compute, model training and inference are becoming major budget line items.",
    "Traditional FinOps tools were built for cloud infrastructure, not AI-specific workloads. They lack token-level visibility, model-aware governance, and AI ROI attribution.",
    "Four critical capability areas define AI FinOps maturity: Telemetry Capture, Quota & Access Enforcement, Metrics (Usage & ROI), and Reporting & Actionable Insights.",
    "Significant whitespace exists, especially in unified platforms that address all four areas for AI workloads. No incumbent covers the full landscape."
]

tb = _add_textbox(slide2, Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.5))
tf = tb.text_frame
tf.word_wrap = True
first = True
for b in bullets:
    if first:
        _set_text(tf, "", font_size=14, color=DARK_GRAY)
        first = False
    p = _add_paragraph(tf, b, font_size=14, color=DARK_GRAY, space_before=Pt(14), space_after=Pt(6))
    p.level = 0

# Accent box at bottom
_add_rect(slide2, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.7), LIGHTEST_BLUE)
tb_note = _add_textbox(slide2, Inches(1.0), Inches(6.3), Inches(11.3), Inches(0.5))
_set_text(tb_note.text_frame,
          "Key Insight: The AI FinOps market is where cloud FinOps was in 2018 \u2014 fragmented, immature, and ripe for a platform play.",
          font_size=12, color=NAVY, bold=True, alignment=PP_ALIGN.LEFT)


# ===========================================================================
# SLIDE 3 : AI FinOps Framework Overview (2x2 grid)
# ===========================================================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
_set_slide_bg(slide3, WHITE)
_add_slide_title(slide3, "AI FinOps Framework Overview")

pillars = [
    ("1", "Telemetry Capture", "Token usage, model calls,\nlatency, GPU utilization"),
    ("2", "Quota & Access Enforcement", "Budgets, rate limits,\naccess policies, model governance"),
    ("3", "Metrics \u2013 Usage & ROI", "Cost per query, cost per outcome,\nunit economics, ROI attribution"),
    ("4", "Reporting & Actionable Insights", "Dashboards, alerts, optimization\nrecommendations, forecasting"),
]

positions = [
    (Inches(0.8),  Inches(1.5)),
    (Inches(6.9),  Inches(1.5)),
    (Inches(0.8),  Inches(4.3)),
    (Inches(6.9),  Inches(4.3)),
]

colors = [DARK_BLUE, ACCENT_BLUE, BLUE, MID_BLUE]

for idx, (num, title, desc) in enumerate(pillars):
    x, y = positions[idx]
    box_w, box_h = Inches(5.6), Inches(2.4)
    # Card background
    _add_rect(slide3, x, y, box_w, box_h, LIGHTEST_BLUE, line_color=LIGHT_BLUE)
    # Number circle
    circ = _add_circle(slide3, x + Inches(0.3), y + Inches(0.3), Inches(0.6), colors[idx])
    circ_tf = circ.text_frame
    circ_tf.word_wrap = False
    p = circ_tf.paragraphs[0]
    p.text = num
    p.font.size = Pt(20)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.name = FONT_NAME
    p.alignment = PP_ALIGN.CENTER
    circ_tf.paragraphs[0].space_before = Pt(0)

    # Title
    tb_t = _add_textbox(slide3, x + Inches(1.1), y + Inches(0.3), Inches(4.2), Inches(0.5))
    _set_text(tb_t.text_frame, title, font_size=18, color=NAVY, bold=True)
    # Description
    tb_d = _add_textbox(slide3, x + Inches(1.1), y + Inches(0.9), Inches(4.2), Inches(1.3))
    _set_text(tb_d.text_frame, desc, font_size=13, color=DARK_GRAY)

# Arrow connectors (simple rectangles)
_add_rect(slide3, Inches(6.4), Inches(2.45), Inches(0.5), Inches(0.12), LIGHT_BLUE)
_add_rect(slide3, Inches(6.4), Inches(5.25), Inches(0.5), Inches(0.12), LIGHT_BLUE)
_add_rect(slide3, Inches(6.55), Inches(3.9), Inches(0.12), Inches(0.4), LIGHT_BLUE)


# ===========================================================================
# SLIDE 4 : Telemetry Capture – What It Covers
# ===========================================================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
_set_slide_bg(slide4, WHITE)
_add_slide_title(slide4, "Telemetry Capture \u2013 What It Covers")

items4 = [
    "Token-level tracking (input/output tokens per request)",
    "Model/provider attribution (which model, which provider, which team)",
    "Latency and throughput metrics",
    "GPU/compute utilization for self-hosted models",
    "Embedding and vector DB query tracking",
    "Data pipeline and training run cost capture",
]

tb4 = _add_textbox(slide4, Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.5))
tf4 = tb4.text_frame
tf4.word_wrap = True
_set_text(tf4, "", font_size=14, color=DARK_GRAY)

for item in items4:
    _add_paragraph(tf4, "\u2022  " + item, font_size=14, color=DARK_GRAY, space_before=Pt(10), space_after=Pt(4))

# Decorative sidebar
_add_rect(slide4, Inches(12.6), Inches(1.3), Inches(0.12), Inches(5.0), LIGHT_BLUE)


# ===========================================================================
# SLIDE 5 : Telemetry Capture – Market Assessment
# ===========================================================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
_set_slide_bg(slide5, WHITE)
_add_slide_title(slide5, "Telemetry Capture \u2013 Market Assessment")

headers5 = ["Solution", "Strengths", "Gaps"]
rows5 = [
    ["Helicone", "Open-source LLM proxy, excellent token tracking, provider-agnostic", "No GPU/training cost capture"],
    ["Portkey", "AI gateway with built-in telemetry, multi-provider routing", "Limited self-hosted model support"],
    ["LangSmith (LangChain)", "Deep tracing for LangChain apps, chain-level cost tracking", "Vendor lock-in to LangChain ecosystem"],
    ["Weights & Biases", "Strong experiment tracking, GPU utilization", "Inference cost tracking is weak"],
    ["OpenTelemetry + custom", "Flexible, open standard", "Requires heavy custom engineering for AI-specific metrics"],
    ["Cloud-native (AWS/Azure/GCP)", "Good compute-level telemetry", "No token-level or model-level granularity"],
]

col_widths5 = [Inches(2.5), Inches(5.0), Inches(4.3)]
_build_table(slide5, Inches(0.7), Inches(1.2), Inches(11.8), col_widths5, headers5, rows5, row_height=Inches(0.55))

# Overall assessment
_add_rect(slide5, Inches(0.7), Inches(5.8), Inches(11.8), Inches(0.7), LIGHTEST_BLUE)
tb5n = _add_textbox(slide5, Inches(0.9), Inches(5.85), Inches(11.4), Inches(0.6))
_set_text(tb5n.text_frame,
          "Overall: Market is fragmented. No single solution captures telemetry across training, inference, and API costs holistically.",
          font_size=12, color=NAVY, bold=True)


# ===========================================================================
# SLIDE 6 : Quota & Access Enforcement – What It Covers
# ===========================================================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
_set_slide_bg(slide6, WHITE)
_add_slide_title(slide6, "Quota & Access Enforcement \u2013 What It Covers")

items6 = [
    "Per-team/per-user budget caps",
    "Rate limiting (requests per minute, tokens per day)",
    "Model access policies (who can use GPT-4 vs GPT-3.5)",
    "Provider routing rules (cost-based, performance-based)",
    "Approval workflows for expensive model usage",
    "Compliance and data governance (PII, data residency)",
]

tb6 = _add_textbox(slide6, Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.5))
tf6 = tb6.text_frame
tf6.word_wrap = True
_set_text(tf6, "", font_size=14, color=DARK_GRAY)
for item in items6:
    _add_paragraph(tf6, "\u2022  " + item, font_size=14, color=DARK_GRAY, space_before=Pt(10), space_after=Pt(4))

_add_rect(slide6, Inches(12.6), Inches(1.3), Inches(0.12), Inches(5.0), LIGHT_BLUE)


# ===========================================================================
# SLIDE 7 : Quota & Access Enforcement – Market Assessment
# ===========================================================================
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
_set_slide_bg(slide7, WHITE)
_add_slide_title(slide7, "Quota & Access Enforcement \u2013 Market Assessment")

headers7 = ["Solution", "Strengths", "Gaps"]
rows7 = [
    ["Portkey", "API-level rate limiting, budget alerts", "No approval workflows, limited RBAC"],
    ["Helicone", "Basic rate limiting", "No quota management or access policies"],
    ["Azure OpenAI", "Token-per-minute quotas, RBAC via Azure AD", "Azure-only, no multi-provider support"],
    ["AWS Bedrock", "IAM-based access, model-level permissions", "AWS-only, no budget-based enforcement"],
    ["Kong / Apigee", "Strong rate limiting, authentication", "Not AI-aware, no token-based quotas"],
    ["Custom solutions", "Many enterprises building in-house", "High maintenance, inconsistent implementation"],
]

_build_table(slide7, Inches(0.7), Inches(1.2), Inches(11.8), col_widths5, headers7, rows7, row_height=Inches(0.55))

_add_rect(slide7, Inches(0.7), Inches(5.8), Inches(11.8), Inches(0.85), LIGHTEST_BLUE)
tb7n = _add_textbox(slide7, Inches(0.9), Inches(5.85), Inches(11.4), Inches(0.75))
tf7n = tb7n.text_frame
tf7n.word_wrap = True
_set_text(tf7n, "Overall: This is the WEAKEST area in the market. No solution provides comprehensive AI-specific quota management "
          "with budget enforcement, approval workflows, and multi-provider RBAC.",
          font_size=12, color=NAVY, bold=True)


# ===========================================================================
# SLIDE 8 : Metrics (Usage & ROI) – What It Covers
# ===========================================================================
slide8 = prs.slides.add_slide(prs.slide_layouts[6])
_set_slide_bg(slide8, WHITE)
_add_slide_title(slide8, "Metrics (Usage & ROI) \u2013 What It Covers")

items8 = [
    "Cost per query / cost per conversation",
    "Cost per business outcome (e.g., cost per ticket resolved, cost per document processed)",
    "Model cost comparison (same task, different models)",
    "Unit economics trending over time",
    "ROI attribution by use case / department",
    "Waste identification (unused provisioned capacity, over-provisioned models)",
]

tb8 = _add_textbox(slide8, Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.5))
tf8 = tb8.text_frame
tf8.word_wrap = True
_set_text(tf8, "", font_size=14, color=DARK_GRAY)
for item in items8:
    _add_paragraph(tf8, "\u2022  " + item, font_size=14, color=DARK_GRAY, space_before=Pt(10), space_after=Pt(4))

_add_rect(slide8, Inches(12.6), Inches(1.3), Inches(0.12), Inches(5.0), LIGHT_BLUE)


# ===========================================================================
# SLIDE 9 : Metrics (Usage & ROI) – Market Assessment
# ===========================================================================
slide9 = prs.slides.add_slide(prs.slide_layouts[6])
_set_slide_bg(slide9, WHITE)
_add_slide_title(slide9, "Metrics (Usage & ROI) \u2013 Market Assessment")

headers9 = ["Solution", "Strengths", "Gaps"]
rows9 = [
    ["Kubecost", "Excellent Kubernetes cost allocation", "No AI/LLM-specific unit economics"],
    ["Vantage", "Multi-cloud cost visibility, good dashboards", "No token-level cost attribution, no ROI metrics"],
    ["Apptio / IBM", "Enterprise TBM, IT financial management", "No AI workload awareness, slow to adapt"],
    ["Harness CCM", "Cloud cost management with recommendations", "Limited AI-specific metrics"],
    ["Helicone / Portkey", "Basic cost tracking per request", "No ROI attribution, no business outcome correlation"],
    ["Weights & Biases", "Experiment cost tracking", "No business ROI, training-focused only"],
]

_build_table(slide9, Inches(0.7), Inches(1.2), Inches(11.8), col_widths5, headers9, rows9, row_height=Inches(0.55))

_add_rect(slide9, Inches(0.7), Inches(5.8), Inches(11.8), Inches(0.7), LIGHTEST_BLUE)
tb9n = _add_textbox(slide9, Inches(0.9), Inches(5.85), Inches(11.4), Inches(0.6))
_set_text(tb9n.text_frame,
          "Overall: Major gap in connecting AI costs to business outcomes. No solution provides ROI attribution or unit economics at the AI workload level.",
          font_size=12, color=NAVY, bold=True)


# ===========================================================================
# SLIDE 10 : Reporting & Actionable Insights – What It Covers
# ===========================================================================
slide10 = prs.slides.add_slide(prs.slide_layouts[6])
_set_slide_bg(slide10, WHITE)
_add_slide_title(slide10, "Reporting & Actionable Insights \u2013 What It Covers")

items10 = [
    "Executive dashboards (total AI spend, trends, forecasts)",
    "Team/project-level cost breakdowns",
    "Anomaly detection and alerting",
    "Optimization recommendations (model downsizing, caching, prompt optimization)",
    "Forecasting and budget planning",
    "Chargeback/showback reports",
]

tb10 = _add_textbox(slide10, Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.5))
tf10 = tb10.text_frame
tf10.word_wrap = True
_set_text(tf10, "", font_size=14, color=DARK_GRAY)
for item in items10:
    _add_paragraph(tf10, "\u2022  " + item, font_size=14, color=DARK_GRAY, space_before=Pt(10), space_after=Pt(4))

_add_rect(slide10, Inches(12.6), Inches(1.3), Inches(0.12), Inches(5.0), LIGHT_BLUE)


# ===========================================================================
# SLIDE 11 : Reporting & Actionable Insights – Market Assessment
# ===========================================================================
slide11 = prs.slides.add_slide(prs.slide_layouts[6])
_set_slide_bg(slide11, WHITE)
_add_slide_title(slide11, "Reporting & Actionable Insights \u2013 Market Assessment")

headers11 = ["Solution", "Strengths", "Gaps"]
rows11 = [
    ["Vantage", "Strong dashboards, multi-cloud", "No AI-specific recommendations"],
    ["CloudHealth (VMware)", "Mature reporting, chargeback", "No AI workload visibility"],
    ["Kubecost", "K8s cost reports, alerts", "No AI optimization recommendations"],
    ["Portkey", "Basic usage dashboards", "No forecasting, limited actionable insights"],
    ["Helicone", "Request-level analytics", "No executive reporting, no optimization suggestions"],
    ["Cast AI", "Automated K8s optimization", "Compute-only, no AI/LLM awareness"],
]

_build_table(slide11, Inches(0.7), Inches(1.2), Inches(11.8), col_widths5, headers11, rows11, row_height=Inches(0.55))

_add_rect(slide11, Inches(0.7), Inches(5.8), Inches(11.8), Inches(0.7), LIGHTEST_BLUE)
tb11n = _add_textbox(slide11, Inches(0.9), Inches(5.85), Inches(11.4), Inches(0.6))
tf11n = tb11n.text_frame
tf11n.word_wrap = True
_set_text(tf11n,
          "Overall: Traditional FinOps tools have strong reporting but zero AI awareness. AI-native tools have basic dashboards but lack mature reporting, forecasting, and optimization engines.",
          font_size=12, color=NAVY, bold=True)


# ===========================================================================
# SLIDE 12 : Competitive Landscape Matrix
# ===========================================================================
slide12 = prs.slides.add_slide(prs.slide_layouts[6])
_set_slide_bg(slide12, WHITE)
_add_slide_title(slide12, "Competitive Landscape Matrix")

# Scoring: 1-5 per pillar
# Columns: Solution | Telemetry | Quota & Access | Metrics (ROI) | Reporting | Total
score_headers = ["Solution", "Telemetry", "Quota & Access", "Metrics (ROI)", "Reporting", "Total"]

score_data = [
    ["Helicone",         "4", "1", "2", "2", "9"],
    ["Portkey",          "4", "2", "2", "2", "10"],
    ["LangSmith",        "3", "1", "2", "2", "8"],
    ["Weights & Biases", "3", "1", "2", "1", "7"],
    ["Kubecost",         "1", "1", "3", "3", "8"],
    ["Vantage",          "1", "1", "2", "4", "8"],
    ["Apptio / IBM",     "1", "1", "2", "3", "7"],
    ["Harness CCM",      "1", "1", "2", "3", "7"],
    ["CloudHealth",      "1", "1", "1", "4", "7"],
    ["Cast AI",          "1", "1", "1", "3", "6"],
    ["Azure / AWS Native","2", "3", "1", "2", "8"],
]

score_col_widths = [Inches(2.3), Inches(1.8), Inches(2.0), Inches(2.0), Inches(1.8), Inches(1.1)]
_build_table(slide12, Inches(0.7), Inches(1.15), Inches(11.0), score_col_widths,
             score_headers, score_data, row_height=Inches(0.4))

# Note at bottom
_add_rect(slide12, Inches(0.7), Inches(6.2), Inches(11.8), Inches(0.7), LIGHTEST_BLUE)
tb12n = _add_textbox(slide12, Inches(0.9), Inches(6.25), Inches(11.4), Inches(0.6))
tf12n = tb12n.text_frame
tf12n.word_wrap = True
_set_text(tf12n,
          "Scores: 1 (No capability) to 5 (Best-in-class). No solution scores above 3 in all four areas. The market lacks a unified AI FinOps platform.",
          font_size=12, color=NAVY, bold=True)

# Legend note
tb12_leg = _add_textbox(slide12, Inches(0.7), Inches(6.85), Inches(6), Inches(0.3))
_set_text(tb12_leg.text_frame, "Scale: 1 = No capability  |  3 = Adequate  |  5 = Best-in-class",
          font_size=10, color=DARK_GRAY)


# ===========================================================================
# SLIDE 13 : Whitespace Analysis
# ===========================================================================
slide13 = prs.slides.add_slide(prs.slide_layouts[6])
_set_slide_bg(slide13, WHITE)
_add_slide_title(slide13, "Whitespace Analysis")

whitespace = [
    ("1", "Unified AI Cost Telemetry Layer",
     "Single SDK/proxy that captures token costs, GPU costs, training costs, and API costs across all providers and self-hosted models."),
    ("2", "AI-Native Quota & Governance Engine",
     "Budget-based enforcement with approval workflows, model-level RBAC, and policy-as-code for AI resources. This is the biggest gap."),
    ("3", "AI ROI Attribution Platform",
     "Connect AI costs to business KPIs (tickets resolved, revenue generated, documents processed). No one does this well."),
    ("4", "Intelligent Optimization Engine",
     "AI-powered recommendations for model selection, prompt optimization, caching strategies, and right-sizing."),
    ("5", "AI Chargeback & Showback",
     "Enterprise-grade departmental cost allocation specifically designed for shared AI infrastructure and API costs."),
]

card_colors = [DARK_BLUE, ACCENT_BLUE, BLUE, MID_BLUE, LIGHT_BLUE]
y_start = Inches(1.25)

for i, (num, title, desc) in enumerate(whitespace):
    y = y_start + Inches(i * 1.15)
    # Number badge
    badge = _add_rect(slide13, Inches(0.8), y + Inches(0.05), Inches(0.5), Inches(0.5), card_colors[i])
    badge_tf = badge.text_frame
    badge_tf.word_wrap = False
    bp = badge_tf.paragraphs[0]
    bp.text = num
    bp.font.size = Pt(18)
    bp.font.color.rgb = WHITE if i < 4 else NAVY
    bp.font.bold = True
    bp.font.name = FONT_NAME
    bp.alignment = PP_ALIGN.CENTER
    badge_tf.paragraphs[0].space_before = Pt(0)
    badge.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Title
    tb_ws_t = _add_textbox(slide13, Inches(1.5), y, Inches(4.0), Inches(0.4))
    _set_text(tb_ws_t.text_frame, title, font_size=15, color=NAVY, bold=True)

    # Description
    tb_ws_d = _add_textbox(slide13, Inches(1.5), y + Inches(0.4), Inches(10.8), Inches(0.6))
    _set_text(tb_ws_d.text_frame, desc, font_size=12, color=DARK_GRAY)


# ===========================================================================
# SLIDE 14 : Recommended Whitespace Strategy
# ===========================================================================
slide14 = prs.slides.add_slide(prs.slide_layouts[6])
_set_slide_bg(slide14, WHITE)
_add_slide_title(slide14, "Recommended Whitespace Strategy")

# Core Platform box
core_x, core_y = Inches(4.0), Inches(1.4)
core_w, core_h = Inches(5.3), Inches(0.65)
_add_rect(slide14, core_x, core_y, core_w, core_h, DARK_BLUE)
tb_core = _add_textbox(slide14, core_x + Inches(0.2), core_y + Inches(0.1), core_w - Inches(0.4), Inches(0.45))
_set_text(tb_core.text_frame, "AI FinOps Control Plane", font_size=18, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

# Layer boxes
layers = [
    ("Layer 1 \u2013 Telemetry", "Universal collector (proxy + SDK) for all AI workloads", ACCENT_BLUE),
    ("Layer 2 \u2013 Governance", "Policy engine with quotas, RBAC, approvals", BLUE),
    ("Layer 3 \u2013 Analytics", "Cost allocation, ROI attribution, unit economics", MID_BLUE),
    ("Layer 4 \u2013 Intelligence", "Optimization recommendations, forecasting, anomaly detection", LIGHT_BLUE),
]

for idx, (lbl, desc, clr) in enumerate(layers):
    ly = Inches(2.25) + Inches(idx * 0.95)
    _add_rect(slide14, Inches(1.2), ly, Inches(4.6), Inches(0.8), clr, line_color=clr)
    tb_lbl = _add_textbox(slide14, Inches(1.4), ly + Inches(0.05), Inches(4.2), Inches(0.35))
    text_clr = WHITE if idx < 2 else NAVY
    _set_text(tb_lbl.text_frame, lbl, font_size=13, color=text_clr, bold=True)
    tb_desc = _add_textbox(slide14, Inches(1.4), ly + Inches(0.38), Inches(4.2), Inches(0.35))
    _set_text(tb_desc.text_frame, desc, font_size=10, color=text_clr)

# Integrations section
integ_x = Inches(6.5)
_add_rect(slide14, integ_x, Inches(2.25), Inches(5.8), Inches(3.8), LIGHTEST_BLUE, line_color=LIGHT_BLUE)
tb_int_title = _add_textbox(slide14, integ_x + Inches(0.2), Inches(2.3), Inches(5.4), Inches(0.4))
_set_text(tb_int_title.text_frame, "Integrations", font_size=16, color=NAVY, bold=True)

integration_groups = [
    ("LLM Providers", "OpenAI, Anthropic, Google, Cohere, Mistral"),
    ("Cloud Platforms", "AWS, Azure, GCP"),
    ("Orchestration", "LangChain, LlamaIndex, Semantic Kernel"),
    ("Infrastructure", "Kubernetes, GPU clusters, vLLM, TGI"),
]

for idx, (grp, items) in enumerate(integration_groups):
    gy = Inches(2.85) + Inches(idx * 0.7)
    _add_rect(slide14, integ_x + Inches(0.3), gy, Inches(0.15), Inches(0.45), ACCENT_BLUE)
    tb_grp = _add_textbox(slide14, integ_x + Inches(0.6), gy, Inches(4.8), Inches(0.25))
    _set_text(tb_grp.text_frame, grp, font_size=12, color=NAVY, bold=True)
    tb_items = _add_textbox(slide14, integ_x + Inches(0.6), gy + Inches(0.25), Inches(4.8), Inches(0.25))
    _set_text(tb_items.text_frame, items, font_size=10, color=DARK_GRAY)

# Differentiation callout
_add_rect(slide14, Inches(1.2), Inches(6.4), Inches(11.1), Inches(0.7), LIGHTEST_BLUE)
tb_diff = _add_textbox(slide14, Inches(1.4), Inches(6.45), Inches(10.7), Inches(0.6))
tf_diff = tb_diff.text_frame
tf_diff.word_wrap = True
_set_text(tf_diff,
          "Differentiation: \"The only platform purpose-built for AI workload financial operations across the full lifecycle.\"",
          font_size=13, color=NAVY, bold=True, alignment=PP_ALIGN.CENTER)


# ===========================================================================
# SLIDE 15 : Next Steps
# ===========================================================================
slide15 = prs.slides.add_slide(prs.slide_layouts[6])
_set_slide_bg(slide15, WHITE)
_add_slide_title(slide15, "Next Steps")

next_steps = [
    ("Validate whitespace hypotheses", "Conduct customer interviews (target: 15\u201320 enterprise AI teams)"),
    ("Build MVP", "Focus on Telemetry + Quota Enforcement (biggest pain, weakest market coverage)"),
    ("Design integration architecture", "Prioritize top 5 LLM providers for Day-1 support"),
    ("Develop pricing model", "Usage-based pricing tied to AI spend managed"),
    ("Target launch timeline & GTM strategy", "Define go-to-market approach and launch milestones"),
]

step_colors = [DARK_BLUE, ACCENT_BLUE, BLUE, MID_BLUE, LIGHT_BLUE]

for idx, (title, desc) in enumerate(next_steps):
    y = Inches(1.4) + Inches(idx * 1.1)

    # Step number circle
    circ = _add_circle(slide15, Inches(0.9), y + Inches(0.05), Inches(0.55), step_colors[idx])
    circ_tf = circ.text_frame
    circ_tf.word_wrap = False
    cp = circ_tf.paragraphs[0]
    cp.text = str(idx + 1)
    cp.font.size = Pt(18)
    cp.font.color.rgb = WHITE if idx < 4 else NAVY
    cp.font.bold = True
    cp.font.name = FONT_NAME
    cp.alignment = PP_ALIGN.CENTER
    circ.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Title
    tb_ns_t = _add_textbox(slide15, Inches(1.7), y, Inches(10.5), Inches(0.35))
    _set_text(tb_ns_t.text_frame, title, font_size=16, color=NAVY, bold=True)
    # Description
    tb_ns_d = _add_textbox(slide15, Inches(1.7), y + Inches(0.38), Inches(10.5), Inches(0.35))
    _set_text(tb_ns_d.text_frame, desc, font_size=13, color=DARK_GRAY)

# Accent line at bottom
_add_rect(slide15, Inches(0.6), Inches(7.0), Inches(12.1), Pt(3), DARK_BLUE)


# ===========================================================================
# Save
# ===========================================================================
output_path = "/home/ubuntu/repos/app_timesheet/AI_FinOps_Market_Landscape.pptx"
prs.save(output_path)
print(f"Presentation saved to {output_path}")
print(f"Total slides: {len(prs.slides)}")
